"""Conversions to and from PDF."""

import base64
import html
import io
import mimetypes
import posixpath
import re
import zipfile
from pathlib import Path

import markdown as md
import pymupdf
from fastapi import HTTPException
from lxml import etree
from PIL import Image, ImageOps

from .util import Workdir, ext, fail, heavy, libreoffice_convert, run, stem

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:  # pragma: no cover - optional
    pass

IMAGE_EXT = {"jpg", "jpeg", "png", "gif", "bmp", "webp", "tif", "tiff", "heic", "heif", "jfif", "ico"}
CHROMIUM_EXT = {"html", "htm", "md", "markdown", "txt", "text", "log", "epub", "xml", "json"}
SELF_URL = "http://127.0.0.1:8080/"

DOC_CSS = """
@page { size: A4; margin: 20mm 18mm; }
body { font-family: 'Noto Sans', 'DejaVu Sans', sans-serif; font-size: 11pt; line-height: 1.5; color: #222; }
h1, h2, h3, h4 { line-height: 1.25; page-break-after: avoid; }
pre, code { font-family: 'DejaVu Sans Mono', monospace; font-size: 9.5pt; }
pre { white-space: pre-wrap; word-break: break-word; background: #f6f6f6; padding: 8px 10px; border-radius: 4px; }
table { border-collapse: collapse; } td, th { border: 1px solid #bbb; padding: 4px 8px; }
img { max-width: 100%; }
blockquote { border-left: 3px solid #e7a; margin-left: 0; padding-left: 12px; color: #555; }
"""


# --------------------------------------------------------------------------- Chromium

def html_to_pdf(html_text: str, *, fmt: str = "A4", landscape: bool = False, margin: str | None = None,
                css_page_size: bool = True, background: bool = True) -> bytes:
    """Render HTML with headless Chromium. Relative URLs resolve against this server (fonts, etc.)."""
    from playwright.sync_api import sync_playwright

    if "<base " not in html_text[:2000].lower():
        html_text = re.sub(r"<head([^>]*)>", rf'<head\1><base href="{SELF_URL}">', html_text, count=1, flags=re.I) \
            if re.search(r"<head", html_text, re.I) else f'<head><base href="{SELF_URL}"></head>' + html_text
    with heavy(), sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
        try:
            page = browser.new_page()
            page.set_content(html_text, wait_until="networkidle", timeout=60000)
            page.evaluate("document.fonts.ready")
            opts = {"print_background": background, "prefer_css_page_size": css_page_size, "landscape": landscape}
            if fmt:
                opts["format"] = fmt
            if margin is not None:
                opts["margin"] = {k: margin for k in ("top", "right", "bottom", "left")}
            return page.pdf(**opts)
        finally:
            browser.close()


# --------------------------------------------------------------------------- to PDF

def images_to_pdf(paths: list[Path]) -> bytes:
    frames = []
    for path in paths:
        with Image.open(path) as im:
            for i in range(getattr(im, "n_frames", 1)):
                im.seek(i)
                frame = ImageOps.exif_transpose(im.copy())
                if frame.mode in ("RGBA", "LA", "P"):
                    bg = Image.new("RGB", frame.size, "white")
                    rgba = frame.convert("RGBA")
                    bg.paste(rgba, mask=rgba.split()[-1])
                    frame = bg
                elif frame.mode not in ("RGB", "L", "CMYK"):
                    frame = frame.convert("RGB")
                frames.append(frame)
    if not frames:
        raise HTTPException(422, "No images found.")
    buf = io.BytesIO()
    frames[0].save(buf, "PDF", save_all=True, append_images=frames[1:], resolution=150)
    return buf.getvalue()


def text_to_html(text: str, title: str) -> str:
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{DOC_CSS} pre{{background:none;padding:0;font-size:10pt}}</style></head><body><pre>{html.escape(text)}</pre></body></html>"


def markdown_to_html(text: str, title: str) -> str:
    body = md.markdown(text, extensions=["extra", "sane_lists", "toc", "nl2br"])
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{DOC_CSS}</style></head><body>{body}</body></html>"


def decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def epub_to_html(path: Path) -> str:
    """Flatten an EPUB's reading order into a single HTML document with inline images."""
    ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container", "opf": "http://www.idpf.org/2007/opf"}
    with zipfile.ZipFile(path) as z:
        container = etree.fromstring(z.read("META-INF/container.xml"))
        opf_path = container.find(".//c:rootfile", ns).get("full-path")
        opf = etree.fromstring(z.read(opf_path))
        base = posixpath.dirname(opf_path)
        manifest = {it.get("id"): it for it in opf.iterfind(".//opf:manifest/opf:item", ns)}
        title = opf.findtext(".//{http://purl.org/dc/elements/1.1/}title") or "EPUB"
        parts = []
        styles = []
        for item in manifest.values():
            if item.get("media-type") == "text/css":
                try:
                    styles.append(z.read(posixpath.join(base, item.get("href"))).decode("utf-8", "replace"))
                except KeyError:
                    pass
        for ref in opf.iterfind(".//opf:spine/opf:itemref", ns):
            item = manifest.get(ref.get("idref"))
            if item is None:
                continue
            doc_path = posixpath.normpath(posixpath.join(base, item.get("href")))
            try:
                tree = etree.fromstring(z.read(doc_path), etree.HTMLParser())
            except (KeyError, etree.XMLSyntaxError):
                continue
            body = tree.find(".//body")
            if body is None:
                continue
            for img in body.iter("img", "{http://www.w3.org/2000/svg}image", "image"):
                attr = "src" if img.get("src") else next((a for a in img.attrib if a.endswith("href")), None)
                if not attr:
                    continue
                src = posixpath.normpath(posixpath.join(posixpath.dirname(doc_path), img.get(attr)))
                try:
                    data = z.read(src)
                except KeyError:
                    continue
                mime = mimetypes.guess_type(src)[0] or "image/jpeg"
                img.set(attr, f"data:{mime};base64,{base64.b64encode(data).decode()}")
            inner = "".join(etree.tostring(ch, encoding="unicode", method="html") for ch in body)
            parts.append(f"<section style='page-break-after:always'>{body.text or ''}{inner}</section>")
    css = DOC_CSS + "\n".join(styles)
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{css}</style></head><body>{''.join(parts)}</body></html>"


def to_pdf(src: Path, wd: Workdir) -> bytes:
    kind = ext(src.name)
    if kind == "pdf":
        return src.read_bytes()
    if kind in IMAGE_EXT:
        return images_to_pdf([src])
    if kind == "svg":
        import cairosvg
        return cairosvg.svg2pdf(url=str(src))
    if kind in ("md", "markdown"):
        return html_to_pdf(markdown_to_html(decode_text(src.read_bytes()), src.stem))
    if kind in ("txt", "text", "log", "json", "xml"):
        return html_to_pdf(text_to_html(decode_text(src.read_bytes()), src.stem))
    if kind in ("html", "htm"):
        return html_to_pdf(decode_text(src.read_bytes()))
    if kind == "epub":
        try:
            return html_to_pdf(epub_to_html(src))
        except (zipfile.BadZipFile, KeyError, AttributeError) as err:
            raise HTTPException(422, "This EPUB file seems to be damaged.") from err
    # Everything else: Office formats (doc, docx, odt, rtf, ppt, pptx, xls, xlsx, ods, odg, pub, csv, …)
    return libreoffice_convert(src, "pdf", wd).read_bytes()


# --------------------------------------------------------------------------- from PDF

def _pdf_to_docx(src: Path, wd: Workdir) -> Path:
    from pdf2docx import Converter
    out = wd.file(f"{src.stem}.docx")
    with heavy():
        cv = Converter(str(src))
        try:
            cv.convert(str(out), multi_processing=False)
        finally:
            cv.close()
    return out


def _pdf_to_xlsx(src: Path, wd: Workdir) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    doc = pymupdf.open(src)
    for page in doc:
        ws = wb.create_sheet(f"Page {page.number + 1}")
        row = 1
        tables = page.find_tables().tables
        if tables:
            for t in tables:
                for r_i, cells in enumerate(t.extract()):
                    for c_i, val in enumerate(cells):
                        cell = ws.cell(row=row, column=c_i + 1, value=_number_or_text(val))
                        if r_i == 0:
                            cell.font = Font(bold=True)
                    row += 1
                row += 1
        else:
            # No tables: one text line per row, words split into columns by large gaps.
            for line in _text_lines(page):
                for c_i, val in enumerate(line):
                    ws.cell(row=row, column=c_i + 1, value=_number_or_text(val))
                row += 1
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    out = wd.file(f"{src.stem}.xlsx")
    wb.save(out)
    return out


def _number_or_text(val):
    if val is None:
        return None
    s = str(val).strip()
    t = s.replace(" ", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*,\d+", t):
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})*(\.\d+)?", t):
        t = t.replace(",", "")
    try:
        return float(t) if "." in t else int(t)
    except ValueError:
        return s


def _text_lines(page):
    words = page.get_text("words")
    lines: dict[tuple, list] = {}
    for w in words:
        lines.setdefault((w[5], w[6]), []).append(w)
    out = []
    for key in sorted(lines, key=lambda k: (min(w[1] for w in lines[k]), k)):
        ws = sorted(lines[key], key=lambda w: w[0])
        cells, cur, last_x1 = [], [], None
        for w in ws:
            if last_x1 is not None and w[0] - last_x1 > 12:
                cells.append(" ".join(cur))
                cur = []
            cur.append(w[4])
            last_x1 = w[2]
        cells.append(" ".join(cur))
        out.append(cells)
    return out


def _pdf_to_epub(src: Path, wd: Workdir) -> Path:
    doc = pymupdf.open(src)
    title = doc.metadata.get("title") or src.stem
    out = wd.file(f"{src.stem}.epub")
    chapters = []
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>""")
        for page in doc:
            body = page.get_text("xhtml")
            name = f"page{page.number + 1:04d}.xhtml"
            chapters.append(name)
            z.writestr(f"OEBPS/{name}", f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Page {page.number + 1}</title>
<style>img{{max-width:100%;height:auto}}</style></head><body>{body}</body></html>""")
        nav = "".join(f'<li><a href="{c}">Page {i + 1}</a></li>' for i, c in enumerate(chapters))
        z.writestr("OEBPS/nav.xhtml", f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head>
<body><nav epub:type="toc"><ol>{nav}</ol></nav></body></html>""")
        items = "".join(f'<item id="p{i}" href="{c}" media-type="application/xhtml+xml"/>' for i, c in enumerate(chapters))
        spine = "".join(f'<itemref idref="p{i}"/>' for i in range(len(chapters)))
        z.writestr("OEBPS/content.opf", f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="uid">urn:maria-pdf-tools:{abs(hash(title))}</dc:identifier>
<dc:title>{html.escape(title)}</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata>
<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>{items}</manifest>
<spine>{spine}</spine></package>""")
    return out


def from_pdf(src: Path, target: str, wd: Workdir) -> tuple[bytes, str, str]:
    """Returns (data, filename, kind)."""
    name = stem(src.name)
    try:
        doc = pymupdf.open(src)
    except Exception as err:
        raise HTTPException(422, "This PDF could not be opened. Try “Repair PDF” first.") from err
    if doc.needs_pass:
        raise HTTPException(422, "This PDF is password protected. Unlock it first ♡")

    if target == "docx":
        return _pdf_to_docx(src, wd).read_bytes(), f"{name}.docx", "docx"
    if target in ("odt", "rtf"):
        docx = _pdf_to_docx(src, wd)
        return libreoffice_convert(docx, target, wd).read_bytes(), f"{name}.{target}", target
    if target == "xlsx":
        return _pdf_to_xlsx(src, wd).read_bytes(), f"{name}.xlsx", "xlsx"
    if target == "ods":
        return libreoffice_convert(_pdf_to_xlsx(src, wd), "ods", wd).read_bytes(), f"{name}.ods", "ods"
    if target in ("pptx", "odp"):
        out = libreoffice_convert(src, target, wd, infilter="impress_pdf_import")
        return out.read_bytes(), f"{name}.{target}", target
    if target == "txt":
        text = "\f".join(page.get_text("text", sort=True) for page in doc)
        return text.encode("utf-8"), f"{name}.txt", "txt"
    if target == "md":
        import pymupdf4llm
        return pymupdf4llm.to_markdown(doc).encode("utf-8"), f"{name}.md", "md"
    if target == "html":
        pages = "".join(f"<div class='page' id='page{p.number + 1}'>{p.get_text('html')}</div>" for p in doc)
        out = f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(doc.metadata.get('title') or name)}</title>
<style>body{{background:#eee;margin:0;padding:20px}} .page{{margin:0 auto 20px;width:fit-content;box-shadow:0 2px 10px rgba(0,0,0,.2);background:#fff}} .page>div{{position:relative}}</style>
</head><body>{pages}</body></html>"""
        return out.encode("utf-8"), f"{name}.html", "html"
    if target == "epub":
        return _pdf_to_epub(src, wd).read_bytes(), f"{name}.epub", "epub"
    if target == "svg":
        if len(doc) == 1:
            return doc[0].get_svg_image(text_as_path=False).encode("utf-8"), f"{name}.svg", "svg"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for p in doc:
                z.writestr(f"{name}-{p.number + 1}.svg", p.get_svg_image(text_as_path=False))
        return buf.getvalue(), f"{name}-svg.zip", "zip"
    if target == "tiff":
        frames = []
        for p in doc:
            pix = p.get_pixmap(dpi=200, alpha=False)
            frames.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        buf = io.BytesIO()
        frames[0].save(buf, "TIFF", save_all=True, append_images=frames[1:], compression="tiff_deflate", dpi=(200, 200))
        return buf.getvalue(), f"{name}.tiff", "tiff"
    raise HTTPException(400, f"Unknown target format: {target}")


# --------------------------------------------------------------------------- images

# Memory a PSD needs, measured with psd-tools: Photoshop's saved preview ≈ the file + ~12 bytes per
# pixel; combining the layers ≈ the file + ~170 bytes per pixel (a 17 MP image: ~3 GB).
PSD_PREVIEW_BYTES_PER_PX = 12
PSD_LAYERS_BYTES_PER_PX = 170


def psd_image(src: Path) -> Image.Image:
    """The whole Photoshop image. Uses the preview Photoshop saves in the file when there is one
    (looks exactly like Photoshop, and is light); the layers are only combined when the preview is
    missing or may lack transparency, and only if the NAS has enough free RAM for it."""
    from psd_tools import PSDImage

    from .util import free_ram_mb

    try:
        psd = PSDImage.open(src)
    except Exception as err:
        raise HTTPException(422, f"{src.name} could not be read as a Photoshop file: {err}") from err
    pixels = psd.width * psd.height
    depth = max(1, psd.depth // 8)
    free = free_ram_mb() * 2**20 * 0.8  # keep a margin for DSM and the rest
    fits_preview = src.stat().st_size + pixels * PSD_PREVIEW_BYTES_PER_PX * depth <= free
    fits_layers = src.stat().st_size + pixels * PSD_LAYERS_BYTES_PER_PX * depth <= free
    size = f"{psd.width}×{psd.height} px"

    preview = None
    if psd.has_preview():
        if not fits_preview:
            raise HTTPException(507, f"{src.name} is {size}; converting it needs more RAM than the NAS has free "
                                     f"right now (about {free_ram_mb()} MB). Close other apps on the NAS and try again.")
        try:
            preview = psd.composite()
        except Exception:  # damaged preview: build the picture from the layers below
            preview = None

    # Photoshop stores transparency in its preview, but some other apps don't. Combining all layers
    # is by far the slowest and hungriest step, so only do it when there is no preview, or when the
    # preview has no transparency AND nothing opaque covers the whole canvas (a normal Background
    # layer does). Some layer kinds can't be redrawn perfectly; whenever that fails, or the NAS
    # doesn't have the RAM for it, Photoshop's own preview is the answer.
    def opaque_cover(layer) -> bool:
        x1, y1, x2, y2 = layer.bbox
        return (layer.is_visible() and layer.opacity == 255 and layer.kind == "pixel"
                and x1 <= 0 and y1 <= 0 and x2 >= psd.width and y2 >= psd.height)

    maybe_transparent = preview is not None and "A" not in preview.getbands() and len(psd) and not any(map(opaque_cover, psd))
    if preview is None or maybe_transparent:
        if not fits_layers:
            if preview is not None:
                return preview
            raise HTTPException(507, f"{src.name} ({size}) has no flattened preview, and combining its layers needs "
                                     f"about {round((src.stat().st_size + pixels * PSD_LAYERS_BYTES_PER_PX * depth) / 2**20)} MB "
                                     f"of RAM — the NAS has about {free_ram_mb()} MB free. In Photoshop, save it with "
                                     "“Maximize PSD compatibility” turned on and try again.")
        try:
            layered = psd.composite(ignore_preview=True)
            if preview is None or ("A" in layered.getbands() and layered.getchannel("A").getextrema()[0] < 255):
                return layered
        except Exception as err:
            if preview is None:
                raise HTTPException(422, f"{src.name} has no flattened preview and its layers could not be combined: {err}. "
                                         "In Photoshop, turn on “Maximize PSD compatibility” when saving and try again.") from err
    return preview


def convert_image(src: Path, to: str, quality: int = 90) -> tuple[bytes, str, str]:
    """Convert one image (incl. HEIC/TIFF/PSD). Multi-frame images become a ZIP of frames."""
    to = "jpg" if to in ("jpg", "jpeg") else "png"
    try:
        # Pillow would read a PSD's layers as frames; psd-tools gives the finished picture instead.
        im = psd_image(src) if ext(src.name) in ("psd", "psb") else Image.open(src)
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(422, f"{src.name} is not an image I can read.") from err
    frames = []
    with im:
        for i in range(getattr(im, "n_frames", 1)):
            im.seek(i)
            # A PSD is already a finished picture in memory: no EXIF rotation, so no extra copy.
            frame = im if isinstance(im, Image.Image) and not hasattr(im, "fp") else ImageOps.exif_transpose(im.copy())
            if to == "jpg":
                if frame.mode in ("RGBA", "LA", "P"):
                    rgba = frame.convert("RGBA")
                    bg = Image.new("RGB", frame.size, "white")
                    bg.paste(rgba, mask=rgba.split()[-1])
                    frame = bg
                elif frame.mode != "RGB":
                    frame = frame.convert("RGB")
            elif frame.mode not in ("RGB", "RGBA", "L", "LA"):
                frame = frame.convert("RGBA")
            buf = io.BytesIO()
            if to == "jpg":
                frame.save(buf, "JPEG", quality=quality, optimize=True)
            elif frame.width * frame.height <= 4_000_000:
                frame.save(buf, "PNG", optimize=True)
            else:
                # optimize=True is ~6× slower for only ~8% smaller big images (one CPU core on a NAS).
                frame.save(buf, "PNG", compress_level=6)
            frames.append(buf.getvalue())
    name = stem(src.name)
    if len(frames) == 1:
        return frames[0], f"{name}.{to}", to
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i, data in enumerate(frames):
            z.writestr(f"{name}-{i + 1}.{to}", data)
    return buf.getvalue(), f"{name}.zip", "zip"


# --------------------------------------------------------------------------- image compression

COMPRESSIBLE = {"jpg", "jpeg", "jfif", "png", "webp", "heic", "heif", "gif", "bmp", "tif", "tiff", "svg"}
IMAGE_LEVELS = {"low": 88, "medium": 76, "high": 62, "extreme": 45}  # quality (or colour budget for PNG/GIF/BMP); "low" is lossless where possible


def _load_frames(src: Path) -> tuple[list[Image.Image], dict]:
    try:
        im = Image.open(src)
    except Exception as err:
        raise HTTPException(422, f"{src.name} is not an image I can read.") from err
    with im:
        info = {k: im.info[k] for k in ("duration", "loop", "transparency") if k in im.info}
        frames = []
        for i in range(getattr(im, "n_frames", 1)):
            im.seek(i)
            frames.append(ImageOps.exif_transpose(im.copy()) if i == 0 else im.copy())
    return frames, info


def _palette(frame: Image.Image, colors: int) -> Image.Image:
    if frame.mode in ("RGBA", "LA") or (frame.mode == "P" and "transparency" in frame.info):
        return frame.convert("RGBA").quantize(colors, method=Image.Quantize.FASTOCTREE, dither=Image.Dither.FLOYDSTEINBERG)
    return frame.convert("RGB").quantize(colors, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.FLOYDSTEINBERG)


def _encode(frames: list[Image.Image], info: dict, kind: str, quality: int, scale: float, lossless: bool) -> bytes:
    """Encode in the same format as the input. `quality` is 1-100; lossy formats use it directly,
    palette formats (PNG/GIF/BMP) turn it into a colour budget."""
    if scale < 0.999:
        frames = [f.resize((max(1, round(f.width * scale)), max(1, round(f.height * scale))), Image.Resampling.LANCZOS)
                  for f in frames]
    colors = max(16, min(256, int(16 + (quality - 20) * 3)))
    buf = io.BytesIO()
    first, rest = frames[0], frames[1:]
    if kind in ("jpg", "jpeg", "jfif"):
        img = first
        if img.mode not in ("RGB", "L", "CMYK"):
            rgba = img.convert("RGBA")
            img = Image.new("RGB", img.size, "white")
            img.paste(rgba, mask=rgba.split()[-1])
        img.save(buf, "JPEG", quality=quality, optimize=True, progressive=True)
    elif kind == "webp":
        extra = {"save_all": True, "append_images": rest, **{k: info[k] for k in ("duration", "loop") if k in info}} if rest else {}
        first.save(buf, "WEBP", quality=quality, method=6, lossless=False, **extra)
    elif kind in ("heic", "heif"):
        first.save(buf, "HEIF", quality=quality)
    elif kind == "png":
        img = first if lossless else _palette(first, colors)
        img.save(buf, "PNG", optimize=True)
    elif kind == "gif":
        out = [f if lossless else _palette(f, colors) for f in frames]
        extra = {k: info[k] for k in ("duration", "loop") if k in info}
        out[0].save(buf, "GIF", save_all=bool(out[1:]), append_images=out[1:], optimize=True, **extra)
    elif kind == "bmp":
        # BMP has no real compression; an 8-bit palette makes it about 3× smaller.
        img = first if lossless else _palette(first, colors).convert("P")
        img.save(buf, "BMP")
    elif kind in ("tif", "tiff"):
        if lossless:
            first.save(buf, "TIFF", compression="tiff_adobe_deflate", save_all=bool(rest), append_images=rest)
        else:
            out = [f.convert("RGB") if f.mode not in ("RGB", "L") else f for f in frames]
            out[0].save(buf, "TIFF", compression="jpeg", quality=quality, save_all=bool(out[1:]), append_images=out[1:])
    else:
        raise HTTPException(422, f"I can't compress .{kind} images.")
    return buf.getvalue()


def _minify_svg(raw: bytes) -> bytes:
    text = decode_text(raw)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = re.sub(r"<metadata\b.*?</metadata>", "", text, flags=re.S | re.I)
    text = re.sub(r"<sodipodi:namedview\b.*?(/>|</sodipodi:namedview>)", "", text, flags=re.S)
    text = re.sub(r"\s(?:inkscape|sodipodi):[\w-]+=\"[^\"]*\"", "", text)
    text = re.sub(r"(\d+\.\d{3})\d+", r"\1", text)  # 3 decimals is plenty for drawing coordinates
    text = re.sub(r">\s+<", "><", text).strip()
    return text.encode("utf-8")


def compress_image(src: Path, level: str = "medium", target: int = 0, max_side: int = 0) -> tuple[bytes, str, str, dict]:
    """Make an image smaller in its own format. With `target` (bytes) search for the best quality that fits."""
    kind = ext(src.name)
    if kind not in COMPRESSIBLE:
        raise HTTPException(422, f"I can't compress .{kind} files. Try JPG, PNG, WEBP, HEIC, GIF, BMP, TIFF or SVG.")
    before = src.stat().st_size
    name = f"{stem(src.name)}-compressed.{kind}"

    def result(data: bytes, note: str = ""):
        if len(data) >= before:
            data, note = src.read_bytes(), note or "already-optimal"
        headers = {"X-Original-Size": str(before), "X-New-Size": str(len(data))}
        if note:
            headers["X-Note"] = note
        return data, name, kind, headers

    if kind == "svg":
        data = _minify_svg(src.read_bytes())
        return result(data, "target-not-reached" if target and len(data) > target else "")

    frames, info = _load_frames(src)
    largest = max(frames[0].size)
    base_scale = min(1.0, max_side / largest) if max_side else 1.0

    # Lossless packing first where the format has it: for drawings and screenshots it often wins.
    lossless = _encode(frames, info, kind, 100, base_scale, lossless=True) if kind in ("png", "gif", "tif", "tiff") else None
    if not target:
        quality = IMAGE_LEVELS.get(level, IMAGE_LEVELS["medium"])
        data = _encode(frames, info, kind, quality, base_scale, lossless=level == "low")
        return result(min(filter(None, (data, lossless)), key=len))

    if before <= target and base_scale == 1.0:
        return result(src.read_bytes(), "already-small")
    # 1) Highest quality that fits at full size (binary search), 2) otherwise make the picture smaller.
    best = lossless if lossless and len(lossless) <= target else None
    lo, hi = (30, 92) if best is None else (1, 0)
    while lo <= hi:
        q = (lo + hi) // 2
        data = _encode(frames, info, kind, q, base_scale, lossless=False)
        if len(data) <= target:
            best, lo = data, q + 1
        else:
            hi = q - 1
    if best is None:
        # Shrink the picture: binary search the largest scale that fits at a decent quality.
        lo_s, hi_s = min(base_scale, 64 / largest), base_scale
        smallest = None
        for _ in range(8):
            scale = (lo_s + hi_s) / 2
            data = _encode(frames, info, kind, 60, scale, lossless=False)
            if smallest is None or len(data) < len(smallest):
                smallest = data
            if len(data) <= target:
                best, lo_s = data, scale
            else:
                hi_s = scale
        if best is None:
            return result(smallest or src.read_bytes(), "target-not-reached")
    return result(best)
